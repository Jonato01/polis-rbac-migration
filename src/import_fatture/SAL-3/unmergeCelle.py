import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import sys


def unmerge_and_fill_excel(input_file, output_file=None):
    """
    Rimuove tutte le celle unite da un file Excel e riempie le righe
    con il contenuto delle celle precedentemente unite.

    Args:
        input_file (str): Percorso del file Excel di input
        output_file (str): Percorso del file Excel di output (opzionale)
    """

    if output_file is None:
        # Se non specificato, crea un nuovo file con suffisso "_unmerged"
        if input_file.endswith('.xlsx'):
            output_file = input_file.replace('.xlsx', '_unmerged.xlsx')
        elif input_file.endswith('.xls'):
            output_file = input_file.replace('.xls', '_unmerged.xlsx')
        else:
            output_file = input_file + '_unmerged.xlsx'

    try:
        # Carica il workbook
        workbook = load_workbook(input_file)

        # Processa ogni foglio di lavoro
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            print(f"Processando foglio: {sheet_name}")

            # Trova tutte le celle unite
            merged_ranges = list(worksheet.merged_cells.ranges)

            if not merged_ranges:
                print(f"  Nessuna cella unita trovata nel foglio {sheet_name}")
                continue

            print(f"  Trovate {len(merged_ranges)} aree di celle unite")

            # Per ogni area di celle unite
            for merged_range in merged_ranges:
                # Ottieni il valore dalla cella superiore sinistra
                top_left_cell = worksheet[merged_range.coord.split(':')[0]]
                value = top_left_cell.value

                # Ottieni le coordinate dell'area unita
                min_col = merged_range.min_col
                max_col = merged_range.max_col
                min_row = merged_range.min_row
                max_row = merged_range.max_row

                print(f"    Elaborando area {merged_range.coord} con valore: {value}")

                # Rimuovi l'unione delle celle
                worksheet.unmerge_cells(merged_range.coord)

                # Riempi tutte le celle dell'area con il valore
                for row in range(min_row, max_row + 1):
                    for col in range(min_col, max_col + 1):
                        cell = worksheet.cell(row=row, column=col)
                        cell.value = value

                        # Copia la formattazione dalla cella originale
                        if top_left_cell.font:
                            cell.font = top_left_cell.font
                        if top_left_cell.fill:
                            cell.fill = top_left_cell.fill
                        if top_left_cell.border:
                            cell.border = top_left_cell.border
                        if top_left_cell.alignment:
                            cell.alignment = top_left_cell.alignment
                        if top_left_cell.number_format:
                            cell.number_format = top_left_cell.number_format

        # Salva il file modificato
        workbook.save(output_file)
        print(f"\nFile salvato come: {output_file}")

    except FileNotFoundError:
        print(f"Errore: File {input_file} non trovato")
    except PermissionError:
        print(f"Errore: Permessi insufficienti per accedere al file {input_file}")
    except Exception as e:
        print(f"Errore durante l'elaborazione: {str(e)}")


def unmerge_excel_alternative(input_file, output_file=None):
    """
    Versione alternativa usando pandas per file più semplici
    (potrebbe non preservare tutta la formattazione)
    """

    if output_file is None:
        if input_file.endswith('.xlsx'):
            output_file = input_file.replace('.xlsx', '_unmerged_pandas.xlsx')
        else:
            output_file = input_file + '_unmerged_pandas.xlsx'

    try:
        # Leggi tutti i fogli
        excel_file = pd.ExcelFile(input_file)

        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            for sheet_name in excel_file.sheet_names:
                print(f"Processando foglio: {sheet_name}")

                # Leggi il foglio
                df = pd.read_excel(input_file, sheet_name=sheet_name, header=None)

                # Riempi i valori NaN con il valore precedente (forward fill)
                # prima per righe, poi per colonne
                df_filled = df.fillna(method='ffill', axis=0).fillna(method='ffill', axis=1)

                # Salva nel nuovo file
                df_filled.to_excel(writer, sheet_name=sheet_name,
                                   index=False, header=False)

        print(f"\nFile salvato come: {output_file}")

    except Exception as e:
        print(f"Errore durante l'elaborazione: {str(e)}")


def main():
    """
    Funzione principale con interfaccia da riga di comando
    """
    print("=" * 60)
    print("PROGRAMMA PER RIMUOVERE CELLE UNITE DA FILE EXCEL")
    print("=" * 60)

    if len(sys.argv) > 1:
        input_file = sys.argv[1]
        output_file = sys.argv[2] if len(sys.argv) > 2 else None
    else:
        # Input interattivo
        input_file = input("Inserisci il percorso del file Excel: ").strip()

        use_custom_output = input("Vuoi specificare un nome per il file di output? (s/n): ").strip().lower()
        if use_custom_output in ['s', 'si', 'sì', 'y', 'yes']:
            output_file = input("Inserisci il percorso del file di output: ").strip()
        else:
            output_file = None

    # Chiedi quale metodo usare
    print("\nScegli il metodo:")
    print("1. Metodo completo (preserva formattazione)")
    print("2. Metodo alternativo pandas (più veloce, formattazione limitata)")

    choice = input("Scelta (1/2): ").strip()

    if choice == "2":
        print("\nUsando metodo pandas...")
        unmerge_excel_alternative(input_file, output_file)
    else:
        print("\nUsando metodo completo...")
        unmerge_and_fill_excel(input_file, output_file)

    print("\nElaborazione completata!")


if __name__ == "__main__":
    main()