"""
Classificazione AI dei file Excel tramite Azure AI Foundry (Opus 4.6).

Analizza nomi file, nomi fogli e header di ogni foglio per determinare
il tipo di import (fatture, personale, servizi, gare, rendicontazione).
"""
import json
import logging
import re
from typing import Optional

import openpyxl

from container_app.azure_ai_bridge import call_ai
from container_app.models import ClassificationResult, FileType

logger = logging.getLogger("container_app.file_classifier")

# Prompt di sistema per il classificatore
CLASSIFIER_SYSTEM_MESSAGE = """Sei un agente specializzato nella classificazione di file Excel
per il sistema di import POLIS. Riceverai informazioni su un file Excel (nome, percorso,
nomi dei fogli, header di ogni foglio) e dovrai determinare il tipo di file tra:

1. **fatture** — File con righe dettagliate di fatture singole. Colonne tipiche:
   "Numero fattura del fornitore", "Importo fattura", "Data fattura", "Fornitore",
   "Accordo quadro", "Buono di consegna", "Numero registrazione fattura"
   Il file si chiama tipicamente "Riepilogo analitico spese sostenute"
2. **personale** — Colonne tipiche: "COGNOME", "NOME", "QUALIFICA", "COSTO MEDIO ORARIO",
   "ORE", mesi specifici (es. "ORE MAG 25", "ORE GIU 25")
3. **servizi** — Colonne tipiche: "Descrizione", "Canale" (SPORTELLO/TOTEM),
   "Quantita'", mesi (Settembre, Ottobre...)
4. **gare** — Colonne tipiche: "Avvio Gara", "Oggetto", "Importo base gara",
   "N. Contratti", "Sottoscrizione contratto"
5. **rendicontazione** — File di rendicontazione siti/sportelli. DEVE avere TUTTE queste colonne chiave:
   "Denominazione sito", "Indirizzo", "Avvii", "Conclusioni" (oltre a "Regione", "Provincia", "Comune").
   Il nome file contiene tipicamente "Rendicontazione" (es. "2025.08_Rendicontazione Sportello Unico.xlsx").
   ATTENZIONE: file che hanno "Regione/Provincia/Comune" ma NON hanno "Denominazione sito", "Indirizzo",
   "Avvii", "Conclusioni" NON sono rendicontazione (es. report installazioni, ecc.)
6. **colonnine** — Report colonnine/stazioni di ricarica. Colonne tipiche:
   "REGIONE", "PROVINCIA", "COMUNE", "Nr IdR installate", "Nr IdR allacciate",
   "ID Piattaforma Unica Nazionale", "Nr BdC".
   Il nome file contiene tipicamente "colonnine" (es. "Report colonnine_al 31_12_2025.xlsx").
   Ha "Regione/Provincia/Comune" ma con colonne specifiche di infrastruttura (IdR, BdC, PUN).

IMPORTANTE - File da classificare come **unknown**:
- File di sintesi/riepilogo (es. "Sintesi 9_SAL.xlsx", "Prospetto sintesi spese sostenute")
- File con colonne aggregate come "RENDICONTATO 1 SAL", "IMPORTO APPROVATO", "CONTRIBUTO"
- File "AQ" (Accordi Quadro) che non sono gare con righe dettagliate
- In generale, se un file contiene dati aggregati/riassuntivi e non righe di dettaglio,
  classificalo come "unknown"

In ogni cartella SAL c'e' tipicamente UN solo file per tipo. Se hai dubbi, usa "unknown".

Rispondi SEMPRE e SOLO con un JSON valido nel formato:
{
  "type": "fatture|personale|servizi|gare|rendicontazione|colonnine|unknown",
  "confidence": 0.95,
  "sal_number": 8,
  "reasoning": "Breve spiegazione della classificazione"
}

Per il sal_number, estrailo dal nome file o dal percorso (es. "8_SAL_..." -> 8,
"SAL-7" -> 7). Se non riesci a determinarlo, usa 0.
"""


def _extract_sal_from_path(blob_path: str) -> Optional[int]:
    """
    Tenta di estrarre il numero SAL dal percorso del blob.
    Patterns supportati:
      - "8_SAL_30_04_2025/"
      - "SAL-7/"
      - "SAL 8/"
      - "7° SAL"
    """
    patterns = [
        r"(\d+)[_\s]*SAL",         # "8_SAL", "8 SAL"
        r"SAL[_\s-]*(\d+)",        # "SAL-7", "SAL 8", "SAL_7"
        r"(\d+)°?\s*SAL",          # "7° SAL"
    ]
    for pattern in patterns:
        match = re.search(pattern, blob_path, re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def read_excel_metadata(file_path: str) -> tuple[list[str], dict[str, list[str]]]:
    """
    Legge i nomi dei fogli e le prime righe di header per ogni foglio.
    Usa openpyxl in read_only mode per velocita'.

    Returns:
        (sheet_names, headers_per_sheet)
    """
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        logger.warning("Impossibile aprire '%s' con openpyxl: %s", file_path, e)
        return [], {}

    sheet_names = wb.sheetnames
    headers_per_sheet: dict[str, list[str]] = {}

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        headers: list[str] = []
        # Leggi le prime 10 righe per trovare header
        for row in ws.iter_rows(max_row=10, values_only=True):
            values = [str(v).strip() if v is not None else "" for v in row]
            # Considera come header la prima riga con almeno 3 valori non vuoti
            non_empty = [v for v in values if v]
            if len(non_empty) >= 3 and not headers:
                headers = values
        headers_per_sheet[sheet_name] = headers

    wb.close()
    return sheet_names, headers_per_sheet


def _classify_with_ai(
    blob_path: str,
    sheet_names: list[str],
    headers_per_sheet: dict[str, list[str]],
) -> ClassificationResult:
    """Usa Azure AI per classificare il file."""

    # Prepara il messaggio con le info del file
    headers_info = ""
    for sheet, hdrs in headers_per_sheet.items():
        headers_info += f"\n  Foglio '{sheet}': {hdrs}"

    prompt = (
        f"Classifica questo file Excel:\n"
        f"- Percorso blob: {blob_path}\n"
        f"- Nome file: {blob_path.split('/')[-1]}\n"
        f"- Fogli: {sheet_names}\n"
        f"- Headers per foglio:{headers_info}\n\n"
        f"Rispondi con il JSON di classificazione."
    )

    reply = call_ai(prompt, system_message=CLASSIFIER_SYSTEM_MESSAGE, max_tokens=1024)

    return _parse_classification_response(reply, blob_path)


def _parse_classification_response(
    reply: str, blob_path: str
) -> ClassificationResult:
    """Parsa la risposta JSON dell'AI."""
    # Pulisci markdown
    text = reply.strip()
    if text.startswith("```json"):
        text = text[7:]
    if text.startswith("```"):
        text = text[3:]
    if text.endswith("```"):
        text = text[:-3]
    text = text.strip()

    try:
        data = json.loads(text)
        file_type_str = data.get("type", "unknown").lower()
        try:
            file_type = FileType(file_type_str)
        except ValueError:
            file_type = FileType.UNKNOWN

        sal_number = int(data.get("sal_number", 0))
        # Se l'AI non ha trovato il SAL, proviamo dal path
        if sal_number == 0:
            sal_number = _extract_sal_from_path(blob_path) or 0

        return ClassificationResult(
            file_type=file_type,
            confidence=float(data.get("confidence", 0.0)),
            sal_number=sal_number,
            reasoning=data.get("reasoning", ""),
        )

    except (json.JSONDecodeError, KeyError, TypeError) as e:
        logger.error(
            "Errore parsing risposta classificatore: %s\nRisposta: %s",
            e, reply,
        )
        # Fallback: prova classificazione euristica
        return _classify_heuristic(blob_path)


def _classify_heuristic(blob_path: str) -> ClassificationResult:
    """
    Classificazione di fallback basata su parole chiave nel nome file.
    Usata quando l'AI non risponde correttamente.
    """
    nome = blob_path.lower()
    sal_number = _extract_sal_from_path(blob_path) or 0

    if "fattur" in nome:
        return ClassificationResult(FileType.FATTURE, 0.7, sal_number, "euristica: 'fattur' nel nome")
    if "personal" in nome:
        return ClassificationResult(FileType.PERSONALE, 0.7, sal_number, "euristica: 'personal' nel nome")
    if "serviz" in nome:
        return ClassificationResult(FileType.SERVIZI, 0.7, sal_number, "euristica: 'serviz' nel nome")
    if "gar" in nome and "gare" in nome:
        return ClassificationResult(FileType.GARE, 0.7, sal_number, "euristica: 'gare' nel nome")
    if "rendicont" in nome:
        return ClassificationResult(FileType.RENDICONTAZIONE, 0.7, sal_number, "euristica: 'rendicont' nel nome")
    if "colonnin" in nome:
        return ClassificationResult(FileType.COLONNINE, 0.7, sal_number, "euristica: 'colonnin' nel nome")

    return ClassificationResult(FileType.UNKNOWN, 0.0, sal_number, "nessuna corrispondenza euristica")


def classify_file(
    blob_path: str,
    sheet_names: list[str],
    headers_per_sheet: dict[str, list[str]],
) -> ClassificationResult:
    """
    Classifica un file Excel determinando tipo e numero SAL.

    Tenta prima la classificazione AI, poi il fallback euristico.
    """
    try:
        result = _classify_with_ai(blob_path, sheet_names, headers_per_sheet)
        logger.info(
            "Classificazione '%s': tipo=%s, SAL=%d, confidence=%.2f",
            blob_path,
            result.file_type.value,
            result.sal_number,
            result.confidence,
        )
        return result

    except Exception as e:
        logger.warning(
            "Classificazione AI fallita per '%s': %s. Uso euristica.",
            blob_path, e,
        )
        return _classify_heuristic(blob_path)
