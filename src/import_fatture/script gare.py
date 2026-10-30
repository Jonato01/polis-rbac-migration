import os
import re
import asyncio
import logging
from datetime import datetime, date
from typing import Optional

import mysql.connector
import openpyxl

# ================= CONFIGURAZIONE =================
SOURCE_DIRECTORY = os.getenv("GARE_SOURCE_DIRECTORY", r"C:\gare da importare")

# Credenziali DB solo da variabili d'ambiente (nessun segreto nel sorgente)
DB_HOST = os.getenv("DB_HOST", "")
DB_PORT = int(os.getenv("DB_PORT", "3306"))
DB_USER = os.getenv("DB_USER", "")
DB_PASSWORD = os.getenv("DB_PASSWORD", "")
DB_NAME = os.getenv("DB_NAME", "polis")
# ==================================================


def safe_value(value, max_length: int = None) -> str:
    """Escape e sanitizza un valore per SQL."""
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (date, datetime)):
        return f"'{value.strftime('%Y-%m-%d')}'"
    str_value = str(value)
    if max_length and len(str_value) > max_length:
        str_value = str_value[:max_length]
    str_value = str_value.replace("\\", "\\\\")
    str_value = str_value.replace("'", "\\'")
    str_value = str_value.replace('"', '\\"')
    str_value = str_value.replace("\n", "\\n")
    str_value = str_value.replace("\r", "\\r")
    str_value = str_value.replace("\t", "\\t")
    str_value = str_value.replace("\x00", "")
    return f"'{str_value}'"


def safe_date(value) -> str:
    """Converte un valore in data SQL (YYYY-MM-DD)."""
    if value is None:
        return "NULL"
    if isinstance(value, datetime):
        return f"'{value.strftime('%Y-%m-%d')}'"
    if isinstance(value, date):
        return f"'{value.strftime('%Y-%m-%d')}'"
    str_val = str(value).strip()
    if not str_val:
        return "NULL"
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y%m%d'):
        try:
            parsed = datetime.strptime(str_val, fmt)
            return f"'{parsed.strftime('%Y-%m-%d')}'"
        except ValueError:
            continue
    return "NULL"

# Setup Logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

HEADER_ROW = 4
FIRST_DATA_ROW = 5


class BatchGareImporter:
    """
    Versione adattata del GareExcelImporter per esecuzione batch locale.
    """

    def __init__(self):
        self.conn = mysql.connector.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            database=DB_NAME,
            charset='utf8mb4',
            collation='utf8mb4_unicode_ci',
            ssl_disabled=False,
        )
        self.conn.autocommit = True

    def parse_excel(self, file_path: str) -> list[dict]:
        """
        Legge il file Excel dal disco e processa le righe.
        """
        logger.info(f"Apertura file: {file_path}")

        # Carica il workbook
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
        except Exception as e:
            logger.error(f"Impossibile aprire il file {file_path}: {e}")
            return []

        # Rileva colonna di partenza cercando "Avvio Gara" nell'header
        col_offset = 0
        for col in range(1, ws.max_column + 1):
            val = ws.cell(HEADER_ROW, col).value
            if val and 'avvio gara' in str(val).lower():
                col_offset = col - 1  # 0 se colonna A, 1 se colonna B, ecc.
                break
        if col_offset:
            logger.info(f"Colonne shiftate di {col_offset} (Avvio Gara in colonna {col_offset + 1})")

        rows = []
        last_row_values = {}

        for row_num in range(FIRST_DATA_ROW, ws.max_row + 1):
            cell_a = self._get_merged_cell_value(ws, row_num, 1 + col_offset)  # Avvio Gara
            cell_b = self._get_merged_cell_value(ws, row_num, 2 + col_offset)  # Oggetto
            cell_c = self._get_merged_cell_value(ws, row_num, 3 + col_offset)  # Importo
            cell_d = self._get_merged_cell_value(ws, row_num, 4 + col_offset)  # N. Contratti
            cell_e = self._get_merged_cell_value(ws, row_num, 5 + col_offset)  # Sottoscrizione
            cell_f = self._get_merged_cell_value(ws, row_num, 6 + col_offset)  # Contratti stipulati

            # La colonna A deve contenere una data valida
            if not isinstance(cell_a, datetime):
                # Se la riga ha dati in altre colonne, usa la data precedente
                if any(v is not None for v in [cell_c, cell_d, cell_e, cell_f]):
                    cell_a = last_row_values.get('avvio_gara_raw')
                else:
                    continue

            # Se dopo carry-forward non abbiamo ancora una data, skip
            if cell_a is None:
                continue

            # Carry-forward colonna B (oggetto)
            if not cell_b or (isinstance(cell_b, str) and not cell_b.strip()):
                cell_b = last_row_values.get('oggetto_raw')

            # Salva valori correnti per carry-forward
            last_row_values = {
                'avvio_gara_raw': cell_a,
                'oggetto_raw': cell_b,
            }

            row_data = {
                'avvio_gara': self._parse_date(cell_a),
                'oggetto': str(cell_b).strip() if cell_b else None,
                'importo_base_gara_mln': self._parse_importo(cell_c),
                'n_contratti': self._parse_int(cell_d),
                'sottoscrizione_contratto': self._parse_date(cell_e),
                'contratti_stipulati': self._parse_int(cell_f),
                'riga_excel': row_num,
            }

            # Salta righe di riepilogo/note senza n_contratti
            if row_data['n_contratti'] is None:
                continue

            rows.append(row_data)

        wb.close()
        return rows

    async def import_to_db(self, rows: list[dict], sal_number: int, filename: str):
        """
        Inserisce i dati nel DB.
        NOTA: Usa sal_number anche come id_sal come richiesto.
        """
        if not rows:
            logger.warning(f"Nessuna riga trovata per {filename}")
            return

        # Id SAL uguale al numero SAL come richiesto
        id_sal = sal_number

        # 1. Pulizia preventiva
        try:
            delete_sql = f"DELETE FROM gare_excel WHERE numero_sal = {sal_number} AND nome_file = {safe_value(filename, 255)}"
            cursor = self.conn.cursor()
            cursor.execute(delete_sql)
            cursor.close()
            logger.info(f"Puliti dati precedenti per SAL {sal_number} - File {filename}")
        except Exception as e:
            logger.error(f"Errore pulizia DB: {e}")

        # 2. Inserimento
        success_count = 0
        for row in rows:
            try:
                query = self._build_insert_query(row, sal_number, id_sal, filename)
                cursor = self.conn.cursor()
                cursor.execute(query)
                cursor.close()
                success_count += 1
            except Exception as e:
                logger.error(f"Errore riga {row.get('riga_excel')}: {e}")

        logger.info(f"COMPLETATO: {filename} -> {success_count}/{len(rows)} righe inserite (SAL {sal_number})")

    # --- Helpers di Parsing (identici al tuo script) ---

    def _parse_date(self, value) -> Optional[str]:
        if value is None: return None
        if isinstance(value, datetime): return value.strftime('%Y-%m-%d')
        str_val = str(value).strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y'):
            try:
                return datetime.strptime(str_val, fmt).strftime('%Y-%m-%d')
            except ValueError:
                continue
        return None

    def _parse_importo(self, value) -> Optional[float]:
        if value is None: return None
        if isinstance(value, (int, float)): return float(value)
        str_val = str(value).strip()
        # Rimuovi formule Excel (es. "=70" → "70")
        if str_val.startswith('='):
            simple = re.match(r'^=(\d+[.,]?\d*)$', str_val)
            if simple:
                str_val = simple.group(1)
            else:
                return None  # Formula complessa, non parsabile
        # Rimuovi tutto il testo: parentesi con contenuto, simboli €, lettere
        str_val = re.sub(r'\([^)]*\)', '', str_val)  # (c), (d), (**), ecc.
        str_val = re.sub(r'[€a-zA-Z%*]', '', str_val)  # Simboli e lettere
        str_val = str_val.strip()
        if not str_val:
            return None
        # Gestione separatori decimali europei
        if ',' in str_val and '.' in str_val:
            if str_val.rfind(',') > str_val.rfind('.'):
                str_val = str_val.replace('.', '').replace(',', '.')
            else:
                str_val = str_val.replace(',', '')
        elif ',' in str_val:
            str_val = str_val.replace(',', '.')
        try:
            return float(str_val)
        except ValueError:
            return None

    def _parse_int(self, value) -> Optional[int]:
        if value is None: return None
        try:
            return int(float(value))
        except:
            return None

    def _get_merged_cell_value(self, ws, row, col):
        """Restituisce il valore reale di una cella, risolvendo i merge."""
        cell = ws.cell(row, col)
        if cell.value is not None:
            return cell.value
        # Controlla se la cella è parte di un merge range
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                # Restituisce il valore della cella top-left del merge
                return ws.cell(merged_range.min_row, merged_range.min_col).value
        return None

    def _build_insert_query(self, row, sal_number, id_sal, filename):
        def _num_or_null(val): return str(val) if val is not None else 'NULL'

        # Path fittizio per coerenza con la tabella
        local_path = os.path.join(SOURCE_DIRECTORY, filename)

        return f"""INSERT INTO gare_excel
            (numero_sal, id_sal, avvio_gara, oggetto, importo_base_gara_mln,
             n_contratti, sottoscrizione_contratto, contratti_stipulati,
             riga_excel, path_file, nome_file)
            VALUES (
                {sal_number},
                {_num_or_null(id_sal)},
                {safe_date(row['avvio_gara'])},
                {safe_value(row['oggetto'], 5000)},
                {_num_or_null(row['importo_base_gara_mln'])},
                {_num_or_null(row['n_contratti'])},
                {safe_date(row['sottoscrizione_contratto'])},
                {_num_or_null(row['contratti_stipulati'])},
                {_num_or_null(row.get('riga_excel'))},
                {safe_value(local_path, 500)},
                {safe_value(filename, 255)}
            )"""


def extract_sal_from_filename(filename: str) -> Optional[int]:
    """
    Estrae il numero SAL da stringhe come 'gare 3 sal.xlsx'
    Regex: cerca un numero tra 'gare' e 'sal' (case insensitive)
    """
    # Pattern: "gare" + spazi/trattini + NUMERO + spazi/trattini + "sal"
    pattern = r"gare[\s_-]+(\d+)[\s_-]+sal"
    match = re.search(pattern, filename, re.IGNORECASE)
    if match:
        return int(match.group(1))
    return None


async def main():
    if not os.path.exists(SOURCE_DIRECTORY):
        logger.error(f"Cartella non trovata: {SOURCE_DIRECTORY}")
        return

    importer = BatchGareImporter()

    files = [f for f in os.listdir(SOURCE_DIRECTORY) if f.lower().endswith(('.xls', '.xlsx'))]
    logger.info(f"Trovati {len(files)} file Excel nella cartella.")

    for filename in files:
        sal_number = extract_sal_from_filename(filename)

        if sal_number is None:
            logger.warning(f"SKIPPATO: Impossibile estrarre numero SAL dal file '{filename}'")
            continue

        full_path = os.path.join(SOURCE_DIRECTORY, filename)

        # 1. Parsing
        rows = importer.parse_excel(full_path)

        if rows:
            # 2. Import DB
            await importer.import_to_db(rows, sal_number, filename)
        else:
            logger.warning(f"Nessuna riga valida trovata in {filename}")


if __name__ == "__main__":
    asyncio.run(main())